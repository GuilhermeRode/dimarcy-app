"""Accounts receivable/payable self-check. Run from backend/:  python test_finance.py
Uses a throwaway SQLite file, never dimarcy.db."""
import os
import tempfile
from datetime import date, datetime, timedelta
from decimal import Decimal

db_file = os.path.join(tempfile.mkdtemp(), "test.db")
os.environ.update(DATABASE_URL=f"sqlite:///{db_file}", ADMIN_EMAIL="admin@test.com",
                  ADMIN_PASSWORD="admin-pass-123", RESEND_API_KEY="", CHECK_EMAIL_DNS="false")

from fastapi.testclient import TestClient  # noqa: E402

from app import clock, finance  # noqa: E402
from app.database import SessionLocal  # noqa: E402
from app.main import app  # noqa: E402

# a customer needs every field but the notes; no city lookup or DNS check over the network
from app.routers import customers as _customers_router  # noqa: E402
_customers_router.geocode_and_cache_background = lambda *a: None
FULL = {"state_registration": "ISENTO", "phone": "(54) 99999-8888", "email": "cliente@gmail.com", "city": "Caxias do Sul", "state": "RS", "address": "Rua A, 1"}
from app.models import FinancialEntry  # noqa: E402

TODAY = date(2026, 10, 15)
clock.today_br = lambda: TODAY

# ---- pure rules ----
assert finance.installment_days("À vista") == [0]
assert finance.installment_days("30/60") == [30, 60]
assert finance.installment_days("30/60/90/120/150") == [30, 60, 90, 120, 150]
assert finance.installment_days("28/56") == [28, 56]
assert finance.installment_days("Quinzenal") == [15, 30, 45, 60, 75, 90, 105, 120, 135, 150]
assert finance.installment_days("combinar") == [0] and finance.installment_days(None) == [0]
assert finance.installment_days("30/²") == [0] and finance.installment_days("99999999") == [0]
assert finance.installment_days("30/3650") == [30, 3650]

assert finance.split_cents(10000, 3) == [3334, 3333, 3333]
assert finance.split_cents(5, 10) == [5] + [0] * 9  # callers drop the zero ones
assert sum(finance.split_cents(123457, 7)) == 123457

assert finance.add_months(date(2027, 1, 31), 0) == date(2027, 1, 31)
assert finance.add_months(date(2027, 1, 31), 1) == date(2027, 2, 28)
assert finance.add_months(date(2027, 1, 31), 2) == date(2027, 3, 31)  # from the first date, no drift
assert finance.add_months(date(2027, 11, 30), 3) == date(2028, 2, 29)  # leap year
assert finance.add_months(date(2026, 12, 10), 1) == date(2027, 1, 10)

assert finance.to_cents(100.1) == 10010 and finance.to_cents(Decimal("33.34")) == 3334
assert finance.to_decimal(0.1 + 0.2) == Decimal("0.30")

e = FinancialEntry(due_date=TODAY, paid_at=None)
assert finance.entry_status(e, TODAY) == "open"
assert finance.entry_status(e, TODAY + timedelta(days=1)) == "overdue"
e.paid_at = TODAY
assert finance.entry_status(e, TODAY + timedelta(days=1)) == "paid"

def login(c, email, password):
    r = c.post("/api/auth/login", json={"email": email, "password": password})
    assert r.status_code == 200, r.text
    return {"Authorization": "Bearer " + r.json()["access_token"]}


def entries(oid):
    """Installments of an order, straight from the database."""
    with SessionLocal() as db:
        return (db.query(FinancialEntry).filter(FinancialEntry.order_id == oid)
                .order_by(FinancialEntry.paid_at.is_(None), FinancialEntry.installment).all())


with TestClient(app) as c:
    admin = login(c, "admin@test.com", "admin-pass-123")
    c.post("/api/users", headers=admin, json={"name": "Ana", "email": "a@t.com", "password": "seller-pass-123",
                                             "role": "seller"})
    A = login(c, "a@t.com", "seller-pass-123")
    r = c.put("/api/settings", headers=admin, json={"allow_price_override": False, "max_discount_percent": 100,
                                                     "retail_markup_percent": 50, "annual_goal": None})
    assert r.status_code == 200, r.text
    color = c.post("/api/colors", headers=admin, json={"name": "Preto", "hex": "#111111"}).json()["id"]
    product = c.post("/api/products", headers=admin, json={"reference": "R1", "description": "Blusão", "price": 100,
                                                           "color_ids": [color]}).json()["id"]
    cheap = c.post("/api/products", headers=admin, json={"reference": "R2", "description": "Botão", "price": 0.05,
                                                         "color_ids": [color]}).json()["id"]
    cust = c.post("/api/customers", headers=admin, json={"name": "Loja Bela", "document": "1", **FULL}).json()["id"]
    cust2 = c.post("/api/customers", headers=admin, json={"name": "Malharia Sul", "document": "2", **FULL}).json()["id"]

    def body(**kw):
        return {"customer_id": cust, "delivery_date": TODAY.isoformat(), "status": "confirmed",
                "payment_method": "Boleto", "payment_terms": "30/60/90", "discount": 0,
                "items": [{"product_id": product, "color_id": color, "size": "P", "quantity": 1}], **kw}

    def new_order(**kw):
        r = c.post("/api/orders", headers=admin, json=body(**kw))
        assert r.status_code == 201, r.text
        return r.json()["id"]

    def set_status(oid, status):
        r = c.patch(f"/api/orders/{oid}/status", headers=admin, json={"status": status})
        assert r.status_code == 200, r.text

    # ---- order hooks ----
    o1 = new_order()  # R$ 100,00 in 30/60/90
    assert entries(o1) == []
    set_status(o1, "delivered")
    e = entries(o1)
    assert [(e_.amount, e_.due_date, e_.installment, e_.installments) for e_ in e] == [
        (Decimal("33.34"), TODAY + timedelta(days=30), 1, 3), (Decimal("33.33"), TODAY + timedelta(days=60), 2, 3),
        (Decimal("33.33"), TODAY + timedelta(days=90), 3, 3)], e
    assert e[0].description == f"Pedido {o1:05d} — parcela 1/3"
    assert (e[0].kind, e[0].payment_method, e[0].paid_at) == ("receivable", "Boleto", None)
    set_status(o1, "delivered")  # same status again: nothing new
    assert len(entries(o1)) == 3

    oq = new_order(status="delivered", payment_terms="Quinzenal",
                   items=[{"product_id": product, "color_id": color, "size": "P", "quantity": 10}])
    assert [x.due_date for x in entries(oq)] == [TODAY + timedelta(days=d) for d in range(15, 151, 15)]
    assert sum(x.amount for x in entries(oq)) == Decimal("1000.00")
    ov = new_order(status="delivered", payment_terms="À vista")
    assert [(x.amount, x.due_date) for x in entries(ov)] == [(Decimal("100.00"), TODAY)]
    oz = new_order(status="delivered", discount=100)  # nothing left to receive
    assert entries(oz) == []
    # Review Focus 1: R$ 0,05 over 10 installments → a single R$ 0,05 entry, never R$ 0,00 ones
    tiny = new_order(status="delivered", payment_terms="Quinzenal",
                     items=[{"product_id": cheap, "color_id": color, "size": "P", "quantity": 1}])
    assert [(x.amount, x.installment, x.installments) for x in entries(tiny)] == [(Decimal("0.05"), 1, 1)]

    # leaving "delivered" drops the open ones and keeps the paid; coming back generates only the balance
    with SessionLocal() as db:
        first = db.query(FinancialEntry).filter_by(order_id=o1, installment=1).one()
        first.paid_at, first.paid_amount = TODAY, first.amount
        db.commit()
    set_status(o1, "shipped")
    assert [(x.installment, x.paid_at) for x in entries(o1)] == [(1, TODAY)]
    set_status(o1, "delivered")
    open_ = [x for x in entries(o1) if x.paid_at is None]
    assert len(open_) == 3 and sum(x.amount for x in open_) == Decimal("66.66"), open_
    set_status(o1, "canceled")
    assert [x.paid_at for x in entries(o1)] == [TODAY]

    # the edit form also counts; editing while still delivered doesn't touch the installments
    oe = new_order()
    r = c.put(f"/api/orders/{oe}", headers=admin, json=body(status="delivered", payment_method="Pix", payment_terms="30/60"))
    assert r.status_code == 200, r.text
    assert [(x.amount, x.payment_method) for x in entries(oe)] == [(Decimal("50.00"), "Pix")] * 2
    r = c.put(f"/api/orders/{oe}", headers=admin, json=body(
        status="delivered", payment_method="Pix", payment_terms="30/60",
        items=[{"product_id": product, "color_id": color, "size": "P", "quantity": 2}]))
    assert r.status_code == 200, r.text
    assert [x.amount for x in entries(oe)] == [Decimal("50.00")] * 2
    assert "installments" not in r.json()  # order payloads (also read by sellers) never carry them

    # an order that still has entries can't be deleted (only quotes can, and a paid one stays)
    set_status(o1, "quote")
    r = c.delete(f"/api/orders/{o1}", headers=admin)
    assert r.status_code == 400 and "parcelas" in r.json()["detail"], r.text
    assert c.delete(f"/api/orders/{new_order(status='quote')}", headers=admin).status_code == 204

    # ---- /api/finance: admin only ----
    for path in ("/api/finance/entries", "/api/finance/summary?kind=payable"):
        assert c.get(path, headers=A).status_code == 403
    assert c.post("/api/finance/entries", headers=A, json={}).status_code == 403
    assert c.post(f"/api/finance/orders/{oe}/generate", headers=A).status_code == 403

    def post_entry(**kw):
        return c.post("/api/finance/entries", headers=admin, json={
            "kind": "payable", "description": "Aluguel", "amount": 1500, "due_date": "2027-01-31",
            "category": "Aluguel", **kw})

    # monthly repetition, from the first date (no drift)
    r = post_entry(repeat_months=12, supplier="Imobiliária")
    assert r.status_code == 201, r.text
    rent = r.json()
    assert len(rent) == 12 and [x["due_date"] for x in rent[:3]] == ["2027-01-31", "2027-02-28", "2027-03-31"]
    assert (rent[2]["description"], rent[2]["installment"], rent[2]["installments"]) == ("Aluguel (3/12)", 3, 12)
    r = post_entry(repeat_months=12, description="x" * 200)
    assert r.status_code == 201, r.text
    assert all(len(x["description"]) <= 200 for x in r.json()) and r.json()[-1]["description"].endswith("(12/12)")
    for x in r.json():  # keep the summary numbers below untouched
        assert c.delete(f"/api/finance/entries/{x['id']}", headers=admin).status_code == 204
    assert post_entry(category="Viagem").status_code == 400
    assert post_entry(amount=0).status_code == 422
    assert post_entry(repeat_months=61).status_code == 422
    assert post_entry(kind="receivable", order_id=999999).status_code == 400

    bill = post_entry(description="Fio", amount=100, due_date=(TODAY - timedelta(days=1)).isoformat(),
                      category="Matéria-prima").json()[0]
    assert (bill["status"], bill["category"], bill["supplier"]) == ("overdue", "Matéria-prima", None)

    # pay / pay twice / can't change a paid amount / unpay
    def pay(eid, **kw):
        return c.post(f"/api/finance/entries/{eid}/pay", headers=admin, json=kw)

    r = pay(bill["id"], paid_amount=98.5)
    assert r.status_code == 200 and (r.json()["status"], r.json()["paid_at"], r.json()["paid_amount"]) == \
        ("paid", TODAY.isoformat(), 98.5), r.text
    assert pay(bill["id"]).status_code == 400
    upd = {"description": "Fio", "amount": 120, "due_date": bill["due_date"], "category": "Matéria-prima"}
    assert c.put(f"/api/finance/entries/{bill['id']}", headers=admin, json=upd).status_code == 400
    r = c.post(f"/api/finance/entries/{bill['id']}/unpay", headers=admin)
    assert (r.json()["paid_at"], r.json()["paid_amount"]) == (None, None)
    assert c.put(f"/api/finance/entries/{bill['id']}", headers=admin, json=upd).json()["amount"] == 120

    # split: total preserved, the original keeps its due date (Review Focus 5)
    split = f"/api/finance/entries/{bill['id']}/split"
    assert c.post(split, headers=admin, json={"amount": 120, "due_date": "2026-12-01"}).status_code == 400
    r = c.post(split, headers=admin, json={"amount": 20, "due_date": "2026-12-01"})
    assert r.status_code == 201, r.text
    part = r.json()
    assert part["description"] == "Fio (restante)"
    found = c.get("/api/finance/entries", headers=admin,
                  params={"kind": "payable", "status": "all", "search": "fio"}).json()
    assert sorted((x["amount"], x["due_date"]) for x in found) == [(20.0, "2026-12-01"), (100.0, bill["due_date"])]
    pay(part["id"])
    assert c.post(f"/api/finance/entries/{part['id']}/split", headers=admin,
                  json={"amount": 5, "due_date": "2026-12-01"}).status_code == 400

    # cards
    s = c.get("/api/finance/summary", headers=admin, params={"kind": "payable"}).json()
    assert s == {"overdue": {"total": 100.0, "count": 1}, "due_soon": {"total": 0.0, "count": 0},
                 "open": {"total": 18100.0, "count": 13}, "paid_this_month": {"total": 20.0, "count": 1}}, s
    assert [x["id"] for x in c.get("/api/finance/entries", headers=admin,
                                    params={"kind": "payable", "status": "overdue"}).json()] == [bill["id"]]
    assert c.get("/api/finance/entries", headers=admin, params={"kind": "payable", "category": "Aluguel",
                                                                "start": "2027-01-01", "end": "2027-02-28"}).json()[1]["due_date"] == "2027-02-28"

    # receivables show the order's *current* customer (Review Focus 2) and search is case-insensitive (3)
    rec = c.get("/api/finance/entries", headers=admin, params={"kind": "receivable", "order_id": oe}).json()
    assert [x["customer_name"] for x in rec] == ["Loja Bela"] * 2
    assert c.get("/api/finance/entries", headers=admin, params={"kind": "receivable", "status": "all",
                                                                "search": "bela"}).json()
    c.put(f"/api/orders/{oe}", headers=admin, json=body(status="delivered", payment_method="Pix", payment_terms="30/60",
                                                         customer_id=cust2))
    rec = c.get("/api/finance/entries", headers=admin, params={"kind": "receivable", "order_id": oe}).json()
    assert [x["customer_name"] for x in rec] == ["Malharia Sul"] * 2

    # "Gerar parcelas" for orders delivered before this feature
    old = new_order(status="delivered", payment_terms="30/60")
    gen = f"/api/finance/orders/{old}/generate"
    assert c.post(gen, headers=admin).status_code == 400  # already has open ones
    with SessionLocal() as db:
        db.query(FinancialEntry).filter(FinancialEntry.order_id == old).delete()
        db.commit()
    r = c.post(gen, headers=admin)
    assert r.status_code == 201 and len(r.json()) == 2, r.text
    assert c.post(f"/api/finance/orders/{new_order()}/generate", headers=admin).status_code == 400  # not delivered
    assert c.post(f"/api/finance/orders/{oz}/generate", headers=admin).status_code == 400  # nothing owed

    # a manual receivable linked to an order
    r = post_entry(kind="receivable", description="Acerto", amount=10, order_id=oe, category=None)
    assert r.status_code == 201 and (r.json()[0]["order_id"], r.json()[0]["category"]) == (oe, None), r.text

    assert c.delete(f"/api/finance/entries/{part['id']}", headers=admin).status_code == 204
    assert c.delete(f"/api/finance/entries/{part['id']}", headers=admin).status_code == 404

    # import from a spreadsheet: preview, then all-or-nothing
    def upload(name, data, kind="payable", apply=False):
        return c.post("/api/finance/import", headers=admin, params={"kind": kind, "apply": apply}, files={"file": (name, data)})

    bank_csv = ("Extrato de outubro\n"
                "Descrição;Fornecedor;Categoria;Valor (R$);Vencimento;Forma de pagamento;Pago em\n"
                "NF 123;Malharia X;matéria-prima;R$ 1.234,56;05/11/2026;Boleto;\n"
                "Luz;CEEE;;-150,00;2026-10-10;Pix;10/10/2026\n"
                ";;;;;;\n").encode("cp1252")
    r = upload("contas.csv", bank_csv)
    assert r.status_code == 200 and r.json()["errors"] == [], r.text
    nf, luz = r.json()["entries"]
    assert (nf["description"], nf["supplier"], nf["category"], nf["amount"], nf["due_date"], nf["paid_at"]) == \
        ("NF 123", "Malharia X", "Matéria-prima", 1234.56, "2026-11-05", None), nf
    assert (luz["category"], luz["amount"], luz["paid_at"], luz["paid_amount"]) == ("Outros", 150.0, "2026-10-10", 150.0), luz
    before = SessionLocal().query(FinancialEntry).count()
    assert upload("contas.csv", bank_csv, apply=True).json() == {"created": 2}
    with SessionLocal() as db:
        assert db.query(FinancialEntry).count() == before + 2

    bad = b"valor,vencimento,categoria\n10,31/02/2026,\nabc,01/01/2027,Viagem\n"
    r = upload("x.csv", bad)
    assert r.json()["entries"] == [] and len(r.json()["errors"]) == 2 and "Linha 2" in r.json()["errors"][0], r.text
    assert upload("x.csv", bad, apply=True).status_code == 400
    assert upload("x.csv", b"nome,telefone\na,b\n").json()["errors"][0].startswith("Não encontrei o cabeçalho")
    assert upload("x.pdf", b"%PDF").status_code == 400

    from io import BytesIO
    from openpyxl import Workbook
    wb = Workbook()
    wb.active.append(["Cliente", "Valor", "Data de vencimento", "Obs"])
    wb.active.append(["Loja Y", 99.9, datetime(2026, 12, 1), "boleto 1/1"])
    buf = BytesIO()
    wb.save(buf)
    r = upload("receber.xlsx", buf.getvalue(), kind="receivable")
    assert r.json()["errors"] == [] and r.json()["entries"][0] == {
        "kind": "receivable", "amount": 99.9, "due_date": "2026-12-01", "payment_method": None, "notes": "boleto 1/1",
        "paid_at": None, "paid_amount": None, "description": "Loja Y"}, r.text
    assert upload("ruim.xlsx", b"not a zip").status_code == 400

    # a credit card statement becomes one payable for its total; the previous bill's payment is left out
    card_csv = (b"date,title,amount\n2026-10-01,Malharia Z,500.00\n2026-10-03,Posto,120.50\n"
                b"2026-10-05,Estorno Posto,-20.50\n2026-10-06,Pagamento recebido,-900.00\n")
    card = {"card_due": "2026-11-10", "card_name": "Nubank", "category": "Matéria-prima"}
    r = c.post("/api/finance/import", headers=admin, params={"kind": "payable", **card}, files={"file": ("n.csv", card_csv)})
    assert r.json()["errors"] == [] and len(r.json()["items"]) == 3, r.text
    (bill,) = r.json()["entries"]
    assert (bill["description"], bill["amount"], bill["due_date"], bill["supplier"], bill["category"]) == \
        ("Fatura Nubank 11/2026", 600.0, "2026-11-10", "Nubank", "Matéria-prima"), bill
    assert bill["notes"].splitlines()[0] == "01/10  Malharia Z  500,00", bill["notes"]
    r = c.post("/api/finance/import", headers=admin, params={"kind": "payable", "apply": True, **card},
               files={"file": ("n.csv", card_csv)})
    assert r.json() == {"created": 1}, r.text

    # boleto: the linha digitável carries the amount and the due date (BB's example from the Febraban spec)
    bb = "00190.50095 40144.816069 06809.350314 3 37370000000100"
    found = finance.parse_boleto_code(bb, TODAY)
    assert (found["amount"], found["due_date"]) == (Decimal("1.00"), date(2032, 8, 21)), found  # factor 3737, new cycle
    assert finance.parse_boleto_code(bb.replace("816069", "816068"), TODAY) is None  # a mistyped digit
    assert finance._boleto_due(9999, TODAY) == date(2025, 2, 21) and finance._boleto_due(1000, TODAY) == date(2025, 2, 22)
    assert finance._boleto_due(1600, TODAY) == date(2026, 10, 15)
    luz = finance.parse_boleto_code("836200000005 667800481000 180975657313 001589636081", TODAY)  # utility, mod 10
    assert luz["amount"] == Decimal("66.78"), luz
    assert finance.parse_boleto_code("846700000017 435900240209 024050002435 842210108119", TODAY)["amount"] == Decimal("143.59")
    assert finance.find_boleto(f"Nosso número 12345678901234567890123\nPague até\n{bb}\nCNPJ 12.345.678/0001-90", TODAY)["code"] == bb

    r = c.post("/api/finance/boleto", headers=admin, data={"code": bb.replace(" ", "").replace(".", "")})
    assert r.status_code == 200 and (r.json()["amount"], r.json()["due_date"]) == (1.0, "2032-08-21"), r.text
    assert c.post("/api/finance/boleto", headers=admin, data={"code": "123"}).status_code == 400
    from fpdf import FPDF
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    pdf.cell(text="Banco do Brasil  Vencimento 21/08/2032")
    pdf.ln()
    pdf.cell(text=bb)
    r = c.post("/api/finance/boleto", headers=admin, files={"file": ("boleto.pdf", bytes(pdf.output()))})
    assert r.status_code == 200 and r.json()["code"] == bb, r.text
    assert c.post("/api/finance/boleto", headers=admin, files={"file": ("b.pdf", b"nada")}).status_code == 400
    assert c.post("/api/finance/boleto", headers=A, data={"code": bb}).status_code == 403  # admins only

    # who's on the boleto: the name after the label, on the same line or the next one
    assert finance.boleto_parties("Beneficiário: MALHARIA SERRANA LTDA - CNPJ 12.345.678/0001-90\nPagador: Loja da Ana ME CPF 123.456.789-00") \
        == {"payee": "MALHARIA SERRANA LTDA", "payer": "Loja da Ana ME"}
    assert finance.boleto_parties("Recibo do Pagador\nBeneficiário Agência/Código do Beneficiário\nFIOS SUL LTDA 1234/56789-0\n"
                                  "Nome do Pagador\nBOUTIQUE BELLA") == {"payee": "FIOS SUL LTDA", "payer": "BOUTIQUE BELLA"}
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    for line in ("Beneficiario: MALHARIA SERRANA LTDA", bb):
        pdf.cell(text=line)
        pdf.ln()
    r = c.post("/api/finance/boleto", headers=admin, files={"file": ("boleto.pdf", bytes(pdf.output()))})
    assert r.json()["payee"] == "MALHARIA SERRANA LTDA", r.text

    # Pix receipts, in the layouts of a few banks
    nubank = ("Comprovante de transferência\n14 OUT 2026 - 10:47:14\nValor R$ 1.250,00\nTipo de transferência Pix\n"
              "Destino\nNome DI MARCY MALHAS LTDA\nCNPJ 12.345.678/0001-90\nOrigem\nNome Ana Paula Souza\n"
              "Instituição NU PAGAMENTOS\nID da transação: E18236120202610291347s0123456789")
    assert finance.parse_pix(nubank, TODAY) == {"amount": Decimal("1250.00"), "paid_at": date(2026, 10, 14),
                                                "payer": "Ana Paula Souza", "payee": "DI MARCY MALHAS LTDA", "e2e_id": "E18236120202610291347s0123456789"}
    inter = "Pix enviado\nR$ 89,90\n13/10/2026 às 09:12\nQuem pagou\nNome\nBOUTIQUE BELLA LTDA\nCPF/CNPJ ***.456.789-**"
    assert finance.parse_pix(inter, TODAY)["payer"] == "BOUTIQUE BELLA LTDA"
    assert finance.parse_pix(inter, TODAY)["amount"] == Decimal("89.90")
    assert finance.parse_pix("De:\nJoão Silva\nPara:\nDi Marcy\nValor: 10,00", TODAY)["payer"] == "João Silva"
    assert finance.parse_pix("nada aqui", TODAY) is None
    assert c.post("/api/finance/pix", headers=admin, files={"file": ("p.png", b"x")}).status_code == 400
    assert finance.parse_pix("Para:\nMalharia Z\nValor: 10,00", TODAY)["payee"] == "Malharia Z"

    # attachments: PDF/PNG/JPG by content, admins only, gone with the entry
    (e,) = c.post("/api/finance/entries", headers=admin, json={"kind": "payable", "description": "Anexo", "amount": 5,
                                                               "due_date": "2026-11-01", "category": "Outros"}).json()
    url = f"/api/finance/entries/{e['id']}/attachment"
    assert c.post(url, headers=admin, files={"file": ("x.pdf", b"<html>")}).status_code == 400
    assert c.post(url, headers=admin, files={"file": ("x.pdf", b"%PDF-1.4 ok")}).json()["attachment"] is True
    r = c.get(url, headers=admin)
    assert r.content == b"%PDF-1.4 ok" and r.headers["content-type"] == "application/pdf", r.headers
    assert c.get(url, headers=A).status_code == 403
    name = SessionLocal().get(FinancialEntry, e["id"]).attachment
    assert c.delete(f"/api/finance/entries/{e['id']}", headers=admin).status_code == 204
    from app.routers.finance import ATTACH_DIR
    assert not (ATTACH_DIR / name).exists()

print("finance checks OK")
