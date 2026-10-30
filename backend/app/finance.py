"""Accounts receivable/payable rules. No FastAPI here, so they're easy to test."""
import calendar
import csv
import io
import re
import unicodedata
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation

from sqlalchemy.orm import Session

from .models import PAYABLE_CATEGORIES, FinancialEntry, Order

QUINZENAL_DAYS = list(range(15, 151, 15))  # "Quinzenal" = 10 installments, every 15 days


def installment_days(terms: str | None) -> list[int]:
    """Days after delivery for each installment of an order's payment terms."""
    t = (terms or "").strip()
    if t.lower() == "quinzenal":
        return QUINZENAL_DAYS
    parts = [p.strip() for p in t.split("/")]
    if all(p.isascii() and p.isdigit() for p in parts) and all(int(p) <= 3650 for p in parts):
        return [int(p) for p in parts]
    return [0]  # "À vista" and anything unknown: one installment, due on delivery


def split_cents(total_cents: int, n: int) -> list[int]:
    """Equal split; the leftover cents go to the first installment."""
    base, rest = divmod(total_cents, n)
    return [base + rest] + [base] * (n - 1)


def add_months(d: date, n: int) -> date:
    """Same day n months later, capped at the month's last day (31/01 + 1 → 28/02)."""
    m = d.month - 1 + n
    year, month = d.year + m // 12, m % 12 + 1
    return date(year, month, min(d.day, calendar.monthrange(year, month)[1]))


def to_decimal(v) -> Decimal:
    return Decimal(str(v)).quantize(Decimal("0.01"), ROUND_HALF_UP)


def to_cents(v) -> int:
    return int(to_decimal(v) * 100)


def entry_status(e: FinancialEntry, today: date) -> str:
    if e.paid_at:
        return "paid"
    return "overdue" if e.due_date < today else "open"


def open_installments(db: Session, order_id: int):
    return db.query(FinancialEntry).filter(FinancialEntry.order_id == order_id, FinancialEntry.paid_at.is_(None))


def create_order_installments(db: Session, order: Order, today: date) -> list[FinancialEntry]:
    """Receivables for what's still owed on a delivered order. Does nothing if it already has open
    installments or nothing is owed. Adds to the session; the caller commits."""
    if open_installments(db, order.id).first():
        return []
    paid = db.query(FinancialEntry).filter(FinancialEntry.order_id == order.id, FinancialEntry.paid_at.isnot(None))
    balance = to_cents(order.total) - sum(to_cents(e.amount) for e in paid)
    if balance <= 0:
        return []
    days = installment_days(order.payment_terms)
    parts = [(d, c) for d, c in zip(days, split_cents(balance, len(days))) if c > 0]
    n = len(parts)
    entries = [FinancialEntry(kind="receivable", description=f"Pedido {order.id:05d} — parcela {i}/{n}",
                              amount=Decimal(c) / 100, due_date=today + timedelta(days=d),
                              payment_method=order.payment_method, order_id=order.id,
                              installment=i, installments=n)
               for i, (d, c) in enumerate(parts, start=1)]
    db.add_all(entries)
    return entries


def remove_open_installments(db: Session, order: Order) -> None:
    """Paid installments stay; the caller commits."""
    open_installments(db, order.id).delete(synchronize_session=False)


# ---- import from a spreadsheet (CSV or XLSX) ----

# Header keyword → field; the first match wins, so the specific ones go first ("forma de pagamento" isn't a date)
IMPORT_COLUMNS = [
    ("payment_method", ("forma", "meio de pagamento")),
    ("paid_amount", ("valor pago", "valor recebido")),
    ("paid_at", ("pago em", "recebido em", "data de pagamento", "data do pagamento", "data pagamento",
                 "data de recebimento", "data do recebimento", "data recebimento", "pagamento", "recebimento")),
    ("due_date", ("vencimento", "venc", "data", "date")),  # "date"/"title"/"amount": Nubank's export
    ("amount", ("valor", "total", "amount")),
    ("supplier", ("fornecedor", "favorecido", "credor", "cliente", "sacado", "pagador")),
    ("category", ("categoria",)),
    ("notes", ("observ", "obs")),
    ("description", ("descri", "historico", "titulo", "documento", "title", "estabelecimento", "lancamento")),
]
IMPORT_MAX_ROWS = 2000


def _plain(v) -> str:
    """'Descrição ' → 'descricao'"""
    s = unicodedata.normalize("NFKD", str(v or "")).encode("ascii", "ignore").decode()
    return " ".join(s.lower().replace("_", " ").split())


def _column(header) -> str | None:
    h = _plain(header)
    return next((f for f, words in IMPORT_COLUMNS if h and any(h.startswith(w) or f" {w}" in f" {h}" for w in words)), None)


def read_sheet(filename: str, data: bytes) -> list[list]:
    """Rows of cells from a .csv (; , or tab, UTF-8 or Windows-1252) or the first sheet of a .xlsx."""
    name = filename.lower()
    if name.endswith(".xlsx"):
        from openpyxl import load_workbook
        try:
            ws = load_workbook(io.BytesIO(data), read_only=True, data_only=True).worksheets[0]
        except Exception:
            raise ValueError("Não foi possível ler a planilha. Salve como .xlsx ou .csv e tente de novo.")
        return [list(r) for r in ws.iter_rows(values_only=True)]
    if not name.endswith((".csv", ".txt")):
        raise ValueError("Envie um arquivo .csv ou .xlsx.")
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = data.decode("cp1252", errors="replace")  # Excel's "CSV" on Windows
    # not csv.Sniffer: the decimal commas in "1.234,56" fool it; a ";" file is what Excel saves in pt-BR
    sample = text[:4096]
    delimiter = next((d for d in ";\t" if d in sample), ",")
    return list(csv.reader(io.StringIO(text), delimiter=delimiter))


def parse_money(v, signed=False) -> Decimal | None:
    """1234.5, 'R$ 1.234,56', '-150,00' (bank exports sign the debits) → positive Decimal.
    signed=True keeps the sign (a card statement's credits) and allows zero."""
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        d = to_decimal(v)
    else:
        s = re.sub(r"[^\d,.-]", "", str(v or ""))
        if "," in s:
            s = s.replace(".", "").replace(",", ".")
        try:
            d = to_decimal(Decimal(s))
        except (InvalidOperation, ValueError):
            return None
    if signed:
        return d if abs(d) <= Decimal("9999999999") else None
    return abs(d) if Decimal(0) < abs(d) <= Decimal("9999999999") else None


def parse_date(v) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v or "").strip().split(" ")[0]
    for fmt in ("%d/%m/%Y", "%d/%m/%y", "%Y-%m-%d", "%d-%m-%Y", "%d.%m.%Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    return None


def _find_header(rows: list[list], required: tuple[str, ...]) -> tuple[int | None, dict]:
    """The first of the top 10 rows naming every required column, so a title above the table is fine."""
    for i, row in enumerate(rows[:10]):
        found = {}
        for j, cell in enumerate(row):
            f = _column(cell)
            if f and f not in found:
                found[f] = j
        if all(f in found for f in required):
            return i, found
    return None, {}


def parse_entries(kind: str, rows: list[list]) -> tuple[list[dict], list[str]]:
    """Spreadsheet rows → (entries ready for FinancialEntry(**e), errors in pt-BR)."""
    header_i, cols = _find_header(rows, ("amount", "due_date"))
    if header_i is None:
        return [], ["Não encontrei o cabeçalho. A planilha precisa ter pelo menos as colunas Valor e Vencimento."]

    categories = {_plain(c): c for c in PAYABLE_CATEGORIES}
    entries, errors = [], []
    for line, row in enumerate(rows[header_i + 1:], start=header_i + 2):
        cell = lambda f: row[cols[f]] if f in cols and cols[f] < len(row) else None  # noqa: E731
        text = lambda f, n: str(cell(f) or "").strip()[:n] or None  # noqa: E731
        if not any(str(c or "").strip() for c in row):
            continue
        problems = []
        amount = parse_money(cell("amount"))
        if amount is None:
            problems.append("valor inválido")
        due = parse_date(cell("due_date"))
        if due is None:
            problems.append("vencimento inválido")
        paid_at = parse_date(cell("paid_at")) if text("paid_at", 40) else None
        if text("paid_at", 40) and paid_at is None:
            problems.append("data de pagamento inválida")
        paid_amount = parse_money(cell("paid_amount")) if text("paid_amount", 40) else None
        if text("paid_amount", 40) and paid_amount is None:
            problems.append("valor pago inválido")
        party = text("supplier", 160)
        description = text("description", 200)
        e = {"kind": kind, "amount": amount, "due_date": due, "payment_method": text("payment_method", 40),
             "notes": text("notes", 5000), "paid_at": paid_at,
             "paid_amount": (paid_amount or amount) if paid_at else None}
        if kind == "payable":
            category = categories.get(_plain(cell("category")) or "outros")
            if category is None:
                problems.append(f"categoria \"{text('category', 40)}\" não existe ({', '.join(PAYABLE_CATEGORIES)})")
            e.update(supplier=party, category=category, description=description or party or "Conta importada")
        else:  # receivables only name a customer through an order, so it goes in the description
            d = " — ".join(x for x in (description, party) if x) or "Recebimento importado"
            e.update(description=d[:200])
        if problems:
            errors.append(f"Linha {line}: {', '.join(problems)}.")
        else:
            entries.append(e)
    if not entries and not errors:
        errors.append("A planilha não tem nenhuma linha preenchida abaixo do cabeçalho.")
    if len(entries) + len(errors) > IMPORT_MAX_ROWS:
        errors.insert(0, f"A planilha tem mais de {IMPORT_MAX_ROWS} linhas. Divida em arquivos menores.")
    return entries, errors


def parse_card_statement(rows: list[list], due: date, card: str, category: str) -> tuple[list[dict], list[str], list[dict]]:
    """A credit card statement → one payable for its total, due on the card's due date, with the purchases
    in the notes. Also returns the purchases, for the preview. The statement's "Data" is the purchase date."""
    header_i, cols = _find_header(rows, ("amount",))
    if header_i is None:
        return [], ["Não encontrei o cabeçalho. A fatura precisa ter pelo menos a coluna Valor."], []
    if category not in PAYABLE_CATEGORIES:
        return [], ["Escolha uma categoria válida."], []
    items, errors = [], []
    for line, row in enumerate(rows[header_i + 1:], start=header_i + 2):
        cell = lambda f: row[cols[f]] if f in cols and cols[f] < len(row) else None  # noqa: E731
        if not any(str(c or "").strip() for c in row):
            continue
        what = " ".join(str(cell(f)).strip() for f in ("description", "supplier") if cell(f)).strip()
        # ponytail: the previous bill's payment shows up as a credit line named "Pagamento ..." in the usual
        # exports (Nubank, Itaú, Inter); it isn't part of this bill. Widen the match if a bank differs.
        if _plain(what).startswith("pagamento"):
            continue
        amount = parse_money(cell("amount"), signed=True)
        if amount is None:
            errors.append(f"Linha {line}: valor inválido.")
            continue
        items.append({"date": parse_date(cell("due_date")), "description": what or "—", "amount": amount})
    total = sum((i["amount"] for i in items), Decimal(0))
    if not errors and not items:
        errors.append("A fatura não tem nenhuma compra abaixo do cabeçalho.")
    if items and not errors and total == 0:
        errors.append("O total da fatura deu zero.")
    if errors:
        return [], errors, items
    lines = (f"{i['date'].strftime('%d/%m') if i['date'] else '--/--'}  {i['description'][:80]}  "
             + f"{i['amount']:.2f}".replace(".", ",") for i in items)
    name = card.strip()[:120] or "Cartão de crédito"
    entry = {"kind": "payable", "description": f"Fatura {name} {due:%m/%Y}", "amount": abs(total), "due_date": due,
             "payment_method": "Cartão", "notes": "\n".join(lines), "paid_at": None, "paid_amount": None,
             "supplier": name, "category": category}  # abs: some banks export the purchases as negatives
    return [entry], [], items


# ---- boleto: the "linha digitável" already carries the amount and the due date (Febraban layout) ----

BOLETO_BASE = date(1997, 10, 7)  # due-date factor 0
BOLETO_RESET = date(2025, 2, 22)  # the factor ran out at 9999 and restarted at 1000 on this day


def _mod10(digits: str) -> int:
    total = 0
    for i, ch in enumerate(reversed(digits)):
        n = int(ch) * (2 if i % 2 == 0 else 1)
        total += n // 10 + n % 10
    return (10 - total % 10) % 10


def _mod11(digits: str, bank: bool) -> int:
    r = sum(int(ch) * (2 + i % 8) for i, ch in enumerate(reversed(digits))) % 11
    if bank:
        return 1 if 11 - r in (0, 10, 11) else 11 - r
    return 0 if r in (0, 1) else 11 - r


def _boleto_due(factor: int, today: date) -> date | None:
    if factor == 0:
        return None  # "contra-apresentação": no due date
    old, new = BOLETO_BASE + timedelta(days=factor), BOLETO_RESET + timedelta(days=factor - 1000)
    return min((old, new), key=lambda d: abs((d - today).days))  # the factor repeats across the reset


def parse_boleto_code(text: str, today: date) -> dict | None:
    """A bank boleto (47 digits) or a utility/tax bill (48 digits, starts with 8) → {amount, due_date, code},
    or None when a check digit doesn't match. amount/due_date are None when the code doesn't carry them."""
    d = re.sub(r"\D", "", text or "")
    if len(d) == 47:
        if any(_mod10(d[a:b]) != int(d[b]) for a, b in ((0, 9), (10, 20), (21, 31))):
            return None
        barcode = d[0:4] + d[32] + d[33:47] + d[4:9] + d[10:20] + d[21:31]
        if _mod11(barcode[:4] + barcode[5:], bank=True) != int(barcode[4]):
            return None
        return {"amount": Decimal(int(d[37:47])) / 100 or None, "due_date": _boleto_due(int(d[33:37]), today),
                "code": f"{d[0:5]}.{d[5:10]} {d[10:15]}.{d[15:21]} {d[21:26]}.{d[26:32]} {d[32]} {d[33:47]}"}
    if len(d) == 48 and d[0] == "8" and d[2] in "6789":
        blocks = [d[i:i + 12] for i in range(0, 48, 12)]
        check = _mod10 if d[2] in "67" else (lambda x: _mod11(x, bank=False))
        if any(check(b[:11]) != int(b[11]) for b in blocks):
            return None
        barcode = "".join(b[:11] for b in blocks)
        amount = Decimal(int(barcode[4:15])) / 100 if d[2] in "68" else None  # 7/9: a reference, not reais
        # ponytail: utility bills have no standard due-date spot; many put YYYYMMDD right after the company code
        try:
            due = datetime.strptime(barcode[19:27], "%Y%m%d").date()
            due = due if abs((due - today).days) <= 730 else None
        except ValueError:
            due = None
        return {"amount": amount or None, "due_date": due, "code": " ".join(blocks)}
    return None


def find_boleto(text: str, today: date) -> dict | None:
    """The first valid linha digitável in a PDF's text; the check digits rule out the other long numbers."""
    for run in re.findall(r"\d[\d.\s-]{42,80}\d", text or ""):
        d = re.sub(r"\D", "", run)
        for size in (48, 47):
            for start in range(len(d) - size + 1):
                found = parse_boleto_code(d[start:start + size], today)
                if found:
                    return found
    return None


# ---- names and Pix receipts: no standard layout, so these are suggestions the user confirms ----

_TAX_ID = re.compile(r"\b(?:CNPJ|CPF)\b.*|\d{2}\.?\d{3}\.?\d{3}/\d{4}-?\d{2}.*|\d{3}\.\d{3}\.\d{3}-\d{2}.*", re.I)
_NOT_NAME = ("agencia", "codigo", "data", "vencimento", "nosso", "valor", "cnpj", "cpf", "endereco", "local", "numero",
             "carteira", "especie", "beneficiario", "pagador", "sacado", "cedente", "recibo", "autentica", "instituicao",
             "banco", "nome", "tipo", "chave", "id ", "identificador", "destino", "origem", "conta", "documento")


def _name(s: str) -> str | None:
    """'MALHARIA X LTDA - CNPJ 12.345.678/0001-90' → 'MALHARIA X LTDA'; None for a label or a code."""
    s = _TAX_ID.sub("", s or "")
    s = re.sub(r"(\s+\S*\d\S*)+$", "", s)  # trailing agency/account codes on the same line
    s = s.strip(" :-–—\t")
    p = _plain(s)
    if len(re.sub(r"[^a-z]", "", p)) < 3 or p.startswith(_NOT_NAME):
        return None
    return s[:160]


def find_name(text: str, label: str) -> str | None:
    """The name after a label that starts a line: on the same line, or else on the next non-empty one."""
    lines = [ln.strip() for ln in (text or "").splitlines()]
    for i, line in enumerate(lines):
        m = re.match(rf"(?:{label})\b\s*:?", line, re.I)
        if not m:
            continue
        name = _name(line[m.end():])
        if not name:
            nxt = next((ln for ln in lines[i + 1:i + 4] if ln and _plain(ln).rstrip(":") != "nome"), "")
            name = _name(re.sub(r"^nome\b\s*:?", "", nxt, flags=re.I))  # "Nome: FULANO" under a section title
        if name:
            return name
    return None


def boleto_parties(text: str) -> dict:
    """Who gets the money (payables) and who pays it (receivables), read from a boleto PDF."""
    return {"payee": find_name(text, r"benefici[aá]rio|cedente"),
            "payer": find_name(text, r"(?:nome do\s+)?(?:pagador|sacado)")}


_MONTHS = ("jan", "fev", "mar", "abr", "mai", "jun", "jul", "ago", "set", "out", "nov", "dez")


def _first_date(text: str, today: date) -> date | None:
    """29/10/2026, '29 OUT 2026' or '29 de outubro de 2026', whichever comes first and is plausible."""
    found = []
    for m in re.finditer(r"\b(\d{1,2})/(\d{1,2})/(\d{4})\b", text):
        found.append((m.start(), int(m[3]), int(m[2]), int(m[1])))
    for m in re.finditer(r"\b(\d{1,2})\s+(?:de\s+)?([a-zç]{3})[a-zç]*\.?\s+(?:de\s+)?(\d{4})\b", text, re.I):
        if m[2].lower() in _MONTHS:
            found.append((m.start(), int(m[3]), _MONTHS.index(m[2].lower()) + 1, int(m[1])))
    for _, y, mo, d in sorted(found):
        try:
            dt = date(y, mo, d)
        except ValueError:
            continue
        if today - timedelta(days=730) <= dt <= today + timedelta(days=1):
            return dt
    return None


def parse_pix(text: str, today: date) -> dict | None:
    """A Pix receipt's text → {amount, paid_at, payer, payee, e2e_id}, or None when there's no amount."""
    m = re.search(r"valor[^\d]{0,40}?(\d{1,3}(?:\.\d{3})*,\d{2})", text or "", re.I) \
        or re.search(r"R\$\s*(\d{1,3}(?:\.\d{3})*,\d{2})", text or "")
    amount = parse_money(m[1]) if m else None
    if amount is None:
        return None
    e2e = re.search(r"\bE\d{20}[0-9A-Za-z]{11}\b", text)  # the end-to-end id every Pix gets from the Central Bank
    return {"amount": amount, "paid_at": _first_date(text, today), "e2e_id": e2e[0] if e2e else None,
            "payer": find_name(text, r"origem|pagador|quem pagou|dados do pagador|dados de quem pagou|de(?=\s*(?::|$))"),
            "payee": find_name(text, r"destino|recebedor|favorecido|quem recebeu|dados do recebedor|dados de quem recebeu"
                                     r"|para(?=\s*(?::|$))")}


def pdf_text(data: bytes) -> str:
    from pypdf import PdfReader
    try:
        return "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(data)).pages[:5])
    except Exception:
        raise ValueError("Não foi possível ler o PDF.")
